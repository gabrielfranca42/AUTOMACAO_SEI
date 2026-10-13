import os
import re
import pandas as pd
import PyPDF2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class PdfAnalyzer:
    def __init__(self):
        self.download_dir = os.path.join(os.getcwd(), "downloads")
        
        # Lista de grupos conhecidos para extração precisa
        self.known_groups = [
            "ABONO PERMANENCIA", "ANOTAÇÃO", "APOSENTADORIA POR INVALIDEZ", "APOSENTADORIA", 
            "AULA ATIVIDADE", "AUXÍLIO FUNERAL", "AVERBAÇÃO", "CARGA HORÁRIA", "CASAMENTO", "CAT", 
            "CERTIDÃO/DECLARAÇÃO", "CREDENCIAMENTO", "CTC", "DECLARAÇÃO DE VÍNCULO", "DEPENDENTE", 
            "DESAVERBAÇÃO", "DESLIGAMENTOS", "DIFÍCIL ACESSO", "ENQUADRAMENTO ADI", 
            "FICHA FINANCEIRA / FICHA FUNCIONAL", "GOZO DE FÉRIAS", "HORA EXTRA", "HORARIO ESPECIAL", 
            "INSALUBRIDADE", "ISENÇÃO DE IMPOSTO DE RENDA", "LDPF", "LICENÇA LUTO", 
            "LICENÇA MATERNIDADE POR ADOÇÃO", "LICENÇA MATERNIDADE", 
            "LICENÇA PARA CASAMENTO", "LICENÇA PARA TRATO DE INTERRESSE PARTICULAR", 
            "LICENÇA PATERNIDADE", "LICENÇA PRÊMIO", "LTS", "MUDANÇA DE NIVEL", "PROGRESSÃO/TITULAÇÃO", 
            "READAPTAÇÃO", "RECONSIDERAÇÃO DE ACRESCIMO", "RECONSIDERAÇÃO DE DESPACHO", "REDUÇÃO DE ATIVIDADES", 
            "REINTEGRAÇÃO", "REQUERIMENTO", "RESTOS DEIXADOS", "REVISÃO DE VENCIMENTOS", "SALÁRIO FAMÍLIA", 
            "VALE ALIMENTAÇÃO"
        ]
        # Ordenar por tamanho decrescente para priorizar nomes mais longos (ex: "LICENÇA MATERNIDADE POR ADOÇÃO" antes de "LICENÇA MATERNIDADE")
        self.known_groups.sort(key=len, reverse=True)

    def _extrair_dados_de_pdfs(self, grupos_alvo=None):
        """Extrai todos os dados dos PDFs encontrados na pasta downloads.
        Retorna lista de dicts com os dados extraídos de cada registro."""
        if not os.path.exists(self.download_dir):
            return [], f"A pasta {self.download_dir} não existe."

        pdf_files = []
        for root, dirs, files in os.walk(self.download_dir):
            for file in files:
                if file.endswith('.pdf'):
                    if grupos_alvo:
                        grupos_alvo_upper = [g.upper() for g in grupos_alvo]
                        rel_path = os.path.relpath(root, self.download_dir).upper()
                        if not any(g in rel_path for g in grupos_alvo_upper):
                            continue
                    pdf_files.append(os.path.join(root, file))

        if not pdf_files:
            return [], "Nenhum arquivo PDF encontrado na pasta downloads."

        extracted_data = []
        for file_path in pdf_files:
            # Identifica o tipo de processo pela pasta do arquivo
            rel_path = os.path.relpath(file_path, self.download_dir)
            pasta_processo = rel_path.split(os.sep)[0] if os.sep in rel_path else ""
            
            try:
                reader = PyPDF2.PdfReader(file_path)
                for page in reader.pages:
                    text = page.extract_text()
                    if not text: continue
                    
                    # Regex resiliente para achar o número do processo (mesmo com espaços em qualquer parte)
                    # Ex: 32.0391 15/2026-64 ou 32.039275/2026-1 1
                    matches = list(re.finditer(r'32\.\s*\d[\d\s]*/\d{4}-\s*\d[\d\s]*\d', text))
                    
                    for i in range(len(matches)):
                        start_pos = matches[i].start()
                        end_pos = matches[i+1].start() if i+1 < len(matches) else len(text)
                        chunk = text[start_pos:end_pos]
                        
                        # Processo limpo sem espaços
                        processo = matches[i].group(0).replace(" ", "")
                        
                        user_date_match = re.search(r'([a-z\._]+)\s+(\d{2}/\d{2}/\d{4})', chunk)
                        usuario = user_date_match.group(1) if user_date_match else ""
                        data = user_date_match.group(2) if user_date_match else ""
                        
                        time_match = re.search(r'(\d{2}:\d{2}:\d{2})', chunk)
                        hora = time_match.group(1) if time_match else ""
                        
                        grupo = ""
                        observacao = ""
                        
                        if time_match:
                            rest_after_time = chunk[time_match.end():].strip()
                            
                            # Encontrar o grupo exato baseado na lista conhecida
                            # Ignora quebras de linha que possam ter partido o nome do grupo no PDF
                            rest_clean = rest_after_time.replace('\n', ' ')
                            matched_group = False
                            for g in self.known_groups:
                                if rest_clean.startswith(g):
                                    grupo = g
                                    observacao = rest_clean[len(g):].strip()
                                    matched_group = True
                                    break
                            
                            if not matched_group:
                                # Se não achou na lista conhecida, extrai a primeira palavra colada ou próxima
                                fallback_match = re.match(r'([A-ZÇÃÕÁÉÍÓÚ]+(?:[\s/][A-ZÇÃÕÁÉÍÓÚ]+)*)', rest_clean)
                                if fallback_match:
                                    grupo = fallback_match.group(1).strip()
                                    observacao = rest_clean[len(grupo):].strip()
                            
                            # Limpar espaços extras na observação
                            observacao = re.sub(r'\s+', ' ', observacao)
                        
                        if processo:
                            extracted_data.append({
                                'Tipo Processo': pasta_processo,
                                'Arquivo Origem': os.path.basename(file_path),
                                'Processo': processo, 'Usuário': usuario,
                                'Data': data, 'Hora': hora, 'Grupo': grupo, 'Observação': observacao
                            })
            except Exception as e: 
                print(f"Erro no PDF {file_path}: {e}")

        return extracted_data, None

    @staticmethod
    def _eh_lts_dobrado(tipo, ano, mes):
        """Verifica se o tipo de processo é LTS e a data é a partir de Dezembro/2025."""
        if str(tipo).strip().upper() != "LTS":
            return False
        try:
            ano_int = int(str(ano).strip())
            mes_int = int(str(mes).strip())
            if ano_int > 2025:
                return True
            elif ano_int == 2025 and mes_int >= 12:
                return True
        except (ValueError, TypeError):
            return False
        return False

    def _gerar_csv_geral(self, df, output_dir):
        """Gera o Excel (resumo_quantitativo.xlsx) no formato de tabela pivô por ano.
        Formato: tipos de processo nas linhas, meses (janeiro-dezembro) nas colunas,
        com contagens e linha TOTAL no final. Uma seção para cada ano.
        Formatação visual profissional com cores, bordas e fontes.
        Regra de negócio: LTS a partir de 12/2025 tem valor dobrado, célula pintada de amarelo e formato 'original (dobrado)'."""
        resumo_path = os.path.join(output_dir, "resumo_quantitativo.xlsx")
        
        # Nomes dos meses em português
        nomes_meses = {
            '01': 'Janeiro', '02': 'Fevereiro', '03': 'Março', '04': 'Abril',
            '05': 'Maio', '06': 'Junho', '07': 'Julho', '08': 'Agosto',
            '09': 'Setembro', '10': 'Outubro', '11': 'Novembro', '12': 'Dezembro'
        }
        meses_ordem = ['01', '02', '03', '04', '05', '06', '07', '08', '09', '10', '11', '12']
        
        total_geral = len(df)
        todos_tipos = sorted(df['Tipo Processo'].unique())
        anos = sorted(df['Ano'].dropna().unique(), reverse=True)
        
        wb = Workbook()
        wb.remove(wb.active)  # Remove a sheet padrão
        
        # ============ ESTILOS ============
        # Bordas
        borda_fina = Side(style='thin', color='B0B0B0')
        borda_media = Side(style='medium', color='4472C4')
        bordas_celula = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
        bordas_header = Border(left=borda_media, right=borda_media, top=borda_media, bottom=borda_media)
        
        # Fontes
        fonte_titulo_ano = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
        fonte_header_mes = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
        fonte_tipo = Font(name='Calibri', size=10, color='1F2937')
        fonte_valor = Font(name='Calibri', size=10, color='374151')
        fonte_total_label = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
        fonte_total_valor = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
        fonte_resumo_titulo = Font(name='Calibri', bold=True, size=16, color='1B3A5C')
        fonte_resumo_info = Font(name='Calibri', size=12, color='374151')
        fonte_resumo_valor = Font(name='Calibri', bold=True, size=12, color='1B3A5C')
        
        # Preenchimentos
        fill_titulo_ano = PatternFill(start_color='1B3A5C', end_color='1B3A5C', fill_type='solid')  # Azul escuro
        fill_header_mes = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')  # Azul médio
        fill_linha_par = PatternFill(start_color='F2F7FB', end_color='F2F7FB', fill_type='solid')    # Azul bem claro
        fill_linha_impar = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')  # Branco
        fill_total = PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid')        # Verde escuro
        fill_lts_dobrado = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')  # Amarelo destaque LTS dobrado
        fill_resumo_bg = PatternFill(start_color='E8F0FE', end_color='E8F0FE', fill_type='solid')    # Azul claro bg
        
        # Alinhamentos
        align_centro = Alignment(horizontal='center', vertical='center')
        align_esquerda = Alignment(horizontal='left', vertical='center')
        align_direita = Alignment(horizontal='right', vertical='center')
        
        # ============ ABA RESUMO GERAL ============
        ws_resumo = wb.create_sheet(title='Resumo Geral')
        
        # Título
        ws_resumo.merge_cells('A1:D1')
        cell_titulo = ws_resumo['A1']
        cell_titulo.value = 'RESUMO GERAL DE PROCESSOS'
        cell_titulo.font = fonte_resumo_titulo
        cell_titulo.alignment = align_esquerda
        
        # Total Secreto (com LTS dobrado a partir de 12/2025)
        total_lts_dobrados = sum(
            1 for _, r in df.iterrows()
            if self._eh_lts_dobrado(r['Tipo Processo'], r.get('Ano', ''), r.get('Mês', ''))
        )
        total_secreto = total_geral + total_lts_dobrados
        
        # Info geral
        ws_resumo['A3'] = 'Total Geral de Registros:'
        ws_resumo['A3'].font = fonte_resumo_info
        if total_secreto != total_geral:
            ws_resumo['B3'] = f"{total_geral} ({total_secreto})"
        else:
            ws_resumo['B3'] = total_geral
        ws_resumo['B3'].font = fonte_resumo_valor
        ws_resumo['B3'].alignment = align_esquerda
        
        ws_resumo['A4'] = 'Tipos de Processo Encontrados:'
        ws_resumo['A4'].font = fonte_resumo_info
        ws_resumo['B4'] = len(todos_tipos)
        ws_resumo['B4'].font = fonte_resumo_valor
        ws_resumo['B4'].alignment = align_esquerda
        
        # Tabela de anos
        ws_resumo['A6'] = 'Ano'
        ws_resumo['B6'] = 'Quantidade'
        for col_cell in ['A6', 'B6']:
            ws_resumo[col_cell].font = fonte_header_mes
            ws_resumo[col_cell].fill = fill_header_mes
            ws_resumo[col_cell].alignment = align_centro
            ws_resumo[col_cell].border = bordas_header
        
        for i, ano in enumerate(anos):
            row = 7 + i
            ws_resumo.cell(row=row, column=1, value=int(ano)).font = fonte_tipo
            ws_resumo.cell(row=row, column=1).alignment = align_centro
            ws_resumo.cell(row=row, column=1).border = bordas_celula
            
            df_ano = df[df['Ano'] == ano]
            qtd_ano_pub = len(df_ano)
            lts_dobrados_ano = sum(
                1 for _, r in df_ano.iterrows()
                if self._eh_lts_dobrado(r['Tipo Processo'], r.get('Ano', ''), r.get('Mês', ''))
            )
            qtd_ano_sec = qtd_ano_pub + lts_dobrados_ano
            
            if qtd_ano_sec != qtd_ano_pub:
                ws_resumo.cell(row=row, column=2, value=f"{qtd_ano_pub} ({qtd_ano_sec})").font = fonte_valor
            else:
                ws_resumo.cell(row=row, column=2, value=qtd_ano_pub).font = fonte_valor
                
            ws_resumo.cell(row=row, column=2).alignment = align_centro
            ws_resumo.cell(row=row, column=2).border = bordas_celula
            if i % 2 == 0:
                ws_resumo.cell(row=row, column=1).fill = fill_linha_par
                ws_resumo.cell(row=row, column=2).fill = fill_linha_par
        
        ws_resumo.column_dimensions['A'].width = 35
        ws_resumo.column_dimensions['B'].width = 18
        
        # ============ UMA ABA POR ANO ============
        for ano in anos:
            df_ano = df[df['Ano'] == ano]
            ws = wb.create_sheet(title=str(int(ano)))
            
            # Agrupa por tipo de processo e mês
            pivot = df_ano.groupby(['Tipo Processo', 'Mês']).size().reset_index(name='Quantidade')
            
            # ---- LINHA 1: Título do ano (mesclado) ----
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
            cell_ano = ws.cell(row=1, column=1)
            cell_ano.value = f'ANO {int(ano)}'
            cell_ano.font = fonte_titulo_ano
            cell_ano.fill = fill_titulo_ano
            cell_ano.alignment = align_centro
            for c in range(1, 14):
                ws.cell(row=1, column=c).fill = fill_titulo_ano
                ws.cell(row=1, column=c).border = bordas_header
            
            # ---- LINHA 2: Cabeçalho dos meses ----
            ws.cell(row=2, column=1, value='Tipo de Processo')
            ws.cell(row=2, column=1).font = fonte_header_mes
            ws.cell(row=2, column=1).fill = fill_header_mes
            ws.cell(row=2, column=1).alignment = align_centro
            ws.cell(row=2, column=1).border = bordas_header
            
            for j, mes_num in enumerate(meses_ordem):
                col = j + 2
                cell_mes = ws.cell(row=2, column=col)
                cell_mes.value = nomes_meses[mes_num]
                cell_mes.font = fonte_header_mes
                cell_mes.fill = fill_header_mes
                cell_mes.alignment = align_centro
                cell_mes.border = bordas_header
            
            # ---- DADOS: Tipos de processo ----
            totais_mes_publico = {m: 0 for m in meses_ordem}
            totais_mes_secreto = {m: 0 for m in meses_ordem}
            
            for i, tipo in enumerate(todos_tipos):
                row = 3 + i
                fill_linha = fill_linha_par if i % 2 == 0 else fill_linha_impar
                
                # Coluna do tipo de processo
                cell_tipo = ws.cell(row=row, column=1, value=tipo)
                cell_tipo.font = fonte_tipo
                cell_tipo.fill = fill_linha
                cell_tipo.alignment = align_esquerda
                cell_tipo.border = bordas_celula
                
                # Colunas dos meses
                for j, mes_num in enumerate(meses_ordem):
                    col = j + 2
                    qtd_filtro = pivot[(pivot['Tipo Processo'] == tipo) & (pivot['Mês'] == mes_num)]
                    
                    cell_val = ws.cell(row=row, column=col)
                    cell_val.alignment = align_centro
                    cell_val.border = bordas_celula
                    
                    if len(qtd_filtro) > 0:
                        qtd = int(qtd_filtro['Quantidade'].values[0])
                        dobrar = self._eh_lts_dobrado(tipo, ano, mes_num)
                        
                        if dobrar:
                            qtd_dobrada = qtd * 2
                            cell_val.value = f"{qtd} ({qtd_dobrada})"
                            cell_val.fill = fill_lts_dobrado
                            cell_val.font = fonte_valor
                            totais_mes_publico[mes_num] += qtd
                            totais_mes_secreto[mes_num] += qtd_dobrada
                        else:
                            cell_val.value = qtd
                            cell_val.fill = fill_linha
                            cell_val.font = fonte_valor
                            totais_mes_publico[mes_num] += qtd
                            totais_mes_secreto[mes_num] += qtd
                    else:
                        cell_val.value = None
                        cell_val.fill = fill_linha
                        cell_val.font = fonte_valor
            
            # ---- LINHA TOTAL ----
            row_total = 3 + len(todos_tipos)
            cell_total_label = ws.cell(row=row_total, column=1, value='TOTAL')
            cell_total_label.font = fonte_total_label
            cell_total_label.fill = fill_total
            cell_total_label.alignment = align_centro
            cell_total_label.border = bordas_header
            
            for j, mes_num in enumerate(meses_ordem):
                col = j + 2
                cell_total = ws.cell(row=row_total, column=col)
                tot_pub = totais_mes_publico.get(mes_num, 0)
                tot_sec = totais_mes_secreto.get(mes_num, 0)
                
                if tot_pub > 0:
                    if tot_sec != tot_pub:
                        cell_total.value = f"{tot_pub} ({tot_sec})"
                    else:
                        cell_total.value = tot_pub
                else:
                    cell_total.value = None
                    
                cell_total.font = fonte_total_valor
                cell_total.fill = fill_total
                cell_total.alignment = align_centro
                cell_total.border = bordas_header
            
            # ---- Ajuste de largura das colunas ----
            ws.column_dimensions['A'].width = 38
            for j in range(len(meses_ordem)):
                col_letter = get_column_letter(j + 2)
                ws.column_dimensions[col_letter].width = 13
            
            # Congelar painéis (fixa cabeçalho)
            ws.freeze_panes = 'B3'
        
        # Salva o arquivo
        wb.save(resumo_path)
        
        return resumo_path

    def _gerar_csvs_por_processo(self, df, output_dir):
        """Gera um CSV individual para cada tipo de processo na pasta 'csv_por_processo'.
        Cada CSV contém: nome do processo, quantitativo e todos os dados extraídos.
        Regra de negócio: se for LTS a partir de 12/2025, o total de registros exibe público e secreto: original (dobrado)."""
        csv_dir = os.path.join(output_dir, "csv_por_processo")
        if not os.path.exists(csv_dir):
            os.makedirs(csv_dir)
        
        tipos = sorted(df['Tipo Processo'].unique())
        arquivos_gerados = []
        
        for tipo in tipos:
            df_tipo = df[df['Tipo Processo'] == tipo].copy()
            total_tipo_pub = len(df_tipo)
            lts_dobrados = sum(
                1 for _, r in df_tipo.iterrows()
                if self._eh_lts_dobrado(tipo, r.get('Ano', ''), r.get('Mês', ''))
            )
            total_tipo_sec = total_tipo_pub + lts_dobrados
            
            # Nome do arquivo sanitizado
            safe_name = tipo.strip().replace("/", "-").replace("\\", "-")
            if not safe_name:
                safe_name = "SEM_TIPO"
            csv_path = os.path.join(csv_dir, f"{safe_name}.csv")
            
            # Escreve o cabeçalho com nome e quantitativo
            with open(csv_path, 'w', encoding='utf-8-sig') as f:
                f.write(f"Tipo de Processo:;{tipo}\n")
                if total_tipo_sec != total_tipo_pub:
                    f.write(f"Total de Registros:;{total_tipo_pub} ({total_tipo_sec})\n")
                else:
                    f.write(f"Total de Registros:;{total_tipo_pub}\n")
                f.write("\n")
            
            # Colunas de dados (exclui 'Tipo Processo' pois já está no cabeçalho)
            colunas_dados = [c for c in df_tipo.columns if c != 'Tipo Processo']
            df_export = df_tipo[colunas_dados]
            
            # Adiciona os dados completos
            df_export.to_csv(csv_path, mode='a', index=False, sep=';', encoding='utf-8-sig')
            
            arquivos_gerados.append(csv_path)
        
        return arquivos_gerados

    def analyze_pdfs_to_csv(self, grupos_alvo=None):
        """Método principal: extrai dados dos PDFs e gera todos os CSVs.
        Retorna (sucesso: bool, mensagem: str)"""
        
        # 1. Extrai dados dos PDFs
        extracted_data, erro = self._extrair_dados_de_pdfs(grupos_alvo)
        
        if erro:
            return False, erro
        
        if not extracted_data:
            return False, "Nenhum dado extraído dos PDFs."
        
        # 2. Cria DataFrame e adiciona colunas de Mês e Ano
        df = pd.DataFrame(extracted_data)
        df['Mês'] = df['Data'].str[3:5]
        df['Ano'] = df['Data'].str[6:10]
        
        # Adiciona coluna de contagem com distinção para LTS dobrado a partir de 12/2025
        df['Contagem'] = df.apply(
            lambda r: "1 (2)" if self._eh_lts_dobrado(r['Tipo Processo'], r['Ano'], r['Mês']) else "1",
            axis=1
        )
        
        output_dir = os.getcwd()
        aviso_erro = ""
        
        # 3. Gera CSV com todos os dados brutos (dados_extraidos.csv)
        try:
            dados_path = os.path.join(output_dir, "dados_extraidos.csv")
            df.to_csv(dados_path, index=False, sep=';', encoding='utf-8-sig')
        except Exception as e:
            aviso_erro += f"\n[!] Erro ao salvar dados_extraidos.csv: {e}"
        
        # 4. Gera CSV geral com resumo por tipo de processo (resumo_quantitativo.xlsx)
        try:
            resumo_path = self._gerar_csv_geral(df, output_dir)
        except Exception as e:
            aviso_erro += f"\n[!] Erro ao salvar resumo_quantitativo.xlsx: {e}"
        
        # 5. Gera CSVs individuais por tipo de processo
        try:
            arquivos_processo = self._gerar_csvs_por_processo(df, output_dir)
        except Exception as e:
            aviso_erro += f"\n[!] Erro ao gerar CSVs por processo: {e}"
            arquivos_processo = []
        
        # 6. Monta mensagem de resumo para o log
        total_geral = len(df)
        total_lts_dobrados = sum(
            1 for _, r in df.iterrows()
            if self._eh_lts_dobrado(r['Tipo Processo'], r['Ano'], r['Mês'])
        )
        total_sec = total_geral + total_lts_dobrados
        tipos = sorted(df['Tipo Processo'].unique())
        
        texto_resumo = f"\nQUANTITATIVOS:\n{'='*40}\n"
        if total_sec != total_geral:
            texto_resumo += f"Total Geral: {total_geral} ({total_sec}) processos\n"
        else:
            texto_resumo += f"Total Geral: {total_geral} processos\n"
        texto_resumo += f"Tipos de Processo: {len(tipos)}\n"
        texto_resumo += f"CSVs individuais gerados: {len(arquivos_processo)}\n"
        texto_resumo += f"\nPor Tipo de Processo:\n{'-'*30}\n"
        
        for tipo in tipos:
            df_tipo = df[df['Tipo Processo'] == tipo]
            lts_dobrados_tipo = sum(
                1 for _, r in df_tipo.iterrows()
                if self._eh_lts_dobrado(tipo, r.get('Ano', ''), r.get('Mês', ''))
            )
            tot_tipo_sec = len(df_tipo) + lts_dobrados_tipo
            
            if tot_tipo_sec != len(df_tipo):
                texto_resumo += f"\n  {tipo}: {len(df_tipo)} ({tot_tipo_sec}) processo(s)\n"
            else:
                texto_resumo += f"\n  {tipo}: {len(df_tipo)} processo(s)\n"
            
            # Mostra resumo por mês/ano
            if 'Ano' in df_tipo.columns and 'Mês' in df_tipo.columns:
                resumo_mes = df_tipo.groupby(['Ano', 'Mês']).size().reset_index(name='Quantidade')
                resumo_mes = resumo_mes.sort_values(by=['Ano', 'Mês'], ascending=[False, False])
                for _, row in resumo_mes.iterrows():
                    qtd_m = row['Quantidade']
                    if self._eh_lts_dobrado(tipo, row['Ano'], row['Mês']):
                        texto_resumo += f"    [{row['Mês']}/{row['Ano']}]: {qtd_m} ({qtd_m * 2}) processo(s)\n"
                    else:
                        texto_resumo += f"    [{row['Mês']}/{row['Ano']}]: {qtd_m} processo(s)\n"
        
        if total_sec != total_geral:
            msg = f"\nExtração concluída! {len(extracted_data)} ({total_sec}) registros processados."
        else:
            msg = f"\nExtração concluída! {len(extracted_data)} registros lidos."
        msg += f"\n\nArquivos gerados:"
        msg += f"\n  - dados_extraidos.csv (todos os dados brutos)"
        msg += f"\n  - resumo_quantitativo.xlsx (resumo geral por tipo)"
        msg += f"\n  - csv_por_processo/ ({len(arquivos_processo)} arquivos individuais)"
        msg += f"{aviso_erro}\n{texto_resumo}"
        
        return True, msg
